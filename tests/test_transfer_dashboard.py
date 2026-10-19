from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pytest
from flask import Flask
from flask_login import LoginManager, UserMixin
from openpyxl import load_workbook
from sqlalchemy import create_engine

from blueprints.transfer_dashboard import bp as transfer_dashboard_blueprint
from services.transfer_dashboard_service import (
    TransferDashboardError,
    build_transfer_dashboard,
    build_transfer_workbook,
    transfer_filters,
    transfer_period,
)


@pytest.fixture()
def transfer_engine():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    with engine.begin() as connection:
        connection.exec_driver_sql("""
            CREATE TABLE RS (
                RSSTAMP TEXT PRIMARY KEY, RESERVA TEXT, ALOJAMENTO TEXT, NOME TEXT,
                ADULTOS INTEGER, CRIANCAS INTEGER, BEBES INTEGER, FEID INTEGER
            )
        """)
        connection.exec_driver_sql("""
            CREATE TABLE TRANSFERS (
                TRANSFERSTAMP TEXT PRIMARY KEY, RSSTAMP TEXT, RESERVA TEXT, ALOJAMENTO TEXT,
                ORIGEM TEXT, DATA_CHEGADA DATE, HORA_CHEGADA TEXT, NUM_VOO TEXT,
                TEM_PARTIDA INTEGER, DATA_PARTIDA DATE, HORA_PARTIDA TEXT, VALOR NUMERIC,
                ESTADO TEXT, DTCRI DATETIME, DTALT DATETIME, NOME_HOSPEDE TEXT, NUM_PESSOAS INTEGER
            )
        """)
        connection.exec_driver_sql("""
            INSERT INTO RS VALUES
            ('rs-1','HMW4H2TD8D','PARAÍSO H','Ana',2,0,0,2),
            ('rs-2','HMZAR4SSEB','Urban Flat','Bruno',2,1,0,2),
            ('rs-other','OUTSIDE','Outro','Fora',1,0,0,3)
        """)
        connection.exec_driver_sql("""
            INSERT INTO TRANSFERS VALUES
            ('tr-1','rs-1','HMW4H2TD8D','PARAÍSO H','airport','2026-10-19','10:30','FR100',0,NULL,NULL,30,'PENDENTE','2026-10-18 10:00:00',NULL,'Ana',2),
            ('tr-2','rs-2','HMZAR4SSEB','Urban Flat','airport','2026-10-21','22:35','TP200',1,'2026-10-25','08:00',60,'ENVIADO','2026-10-18 11:00:00','2026-10-18 11:05:00','Bruno',3),
            ('tr-3','rs-2','HMZAR4SSEB','Urban Flat','to_airport','2026-10-25','08:00','',1,'2026-10-25','08:00',30,'CANCELADO','2026-10-18 12:00:00',NULL,'Bruno',3),
            ('tr-other','rs-other','OUTSIDE','Outro','airport','2026-10-19','12:00','',0,NULL,NULL,30,'PENDENTE','2026-10-18 13:00:00',NULL,'Fora',1)
        """)
    return engine


def test_period_defaults_to_today_and_next_seven_days():
    period = transfer_period(today=date(2026, 10, 19))
    assert period["mode"] == "next7"
    assert period["start"] == date(2026, 10, 19)
    assert period["end"] == date(2026, 10, 26)
    assert transfer_period("all", today=date(2026, 10, 19))["start"] is None
    with pytest.raises(TransferDashboardError):
        transfer_period("custom", "2026-10-20", "2026-10-19")


def test_dashboard_limits_feid_keeps_duplicates_and_counts_filtered_rows(transfer_engine):
    period = transfer_period("next7", today=date(2026, 10, 19))
    payload = build_transfer_dashboard(
        transfer_engine,
        period,
        today=date(2026, 10, 19),
        now=datetime(2026, 10, 19, 9),
    )

    assert payload["kpis"] == {
        "total": 3,
        "pending": 1,
        "sent": 1,
        "cancelled": 1,
        "arrivals_today": 1,
        "departures_today": 0,
    }
    assert {row["transfer_stamp"] for row in payload["rows"]} == {"tr-1", "tr-2", "tr-3"}
    urban = [row for row in payload["rows"] if row["reservation"] == "HMZAR4SSEB"]
    assert len(urban) == 2
    assert all(row["multiple_for_reservation"] for row in urban)
    assert all(row["reservation_transfer_count"] == 2 for row in urban)
    departure = next(row for row in urban if row["transfer_stamp"] == "tr-3")
    assert departure["arrival_date"] is None
    assert departure["departure_date"] == "2026-10-25"


def test_combined_filters_update_rows_and_kpis(transfer_engine):
    period = transfer_period("all", today=date(2026, 10, 19))
    filters = transfer_filters({
        "movement": "departure",
        "accommodation": "Urban Flat",
        "status": "enviado",
        "reservation": "HMZAR",
    })
    payload = build_transfer_dashboard(transfer_engine, period, filters, today=date(2026, 10, 19))

    assert [row["transfer_stamp"] for row in payload["rows"]] == ["tr-2"]
    assert payload["kpis"]["total"] == 1
    assert payload["kpis"]["sent"] == 1
    assert payload["kpis"]["cancelled"] == 0


def test_excel_contains_filtered_rows_and_metadata(transfer_engine):
    payload = build_transfer_dashboard(
        transfer_engine,
        transfer_period("today", today=date(2026, 10, 19)),
        today=date(2026, 10, 19),
    )
    content = build_transfer_workbook(payload, exported_at=datetime(2026, 10, 19, 15, 30))
    workbook = load_workbook(BytesIO(content))
    sheet = workbook["Transfers"]

    assert sheet["A1"].value == "Resumo de pedidos de transfer"
    assert "19/10/2026 15:30" in sheet["A4"].value
    assert sheet["A7"].value == "HMW4H2TD8D"
    assert sheet.max_row == 7


def test_dashboard_page_requires_login_guestspa_entity_and_access():
    app = Flask(__name__, template_folder=str(Path(__file__).resolve().parents[1] / "templates"))
    app.secret_key = "test"
    login_manager = LoginManager(app)
    login_manager.login_view = "login"
    app.add_url_rule("/login", "login", lambda: "login")
    app.register_blueprint(transfer_dashboard_blueprint)

    class User(UserMixin):
        def __init__(self, user_id, admin):
            self.id = user_id
            self.LOGIN = user_id
            self.ADMIN = admin

    users = {"admin": User("admin", True)}

    @login_manager.user_loader
    def load_user(user_id):
        return users.get(user_id)

    client = app.test_client()
    assert client.get("/dashboard/transfers").status_code == 302
    with client.session_transaction() as session:
        session["_user_id"] = "admin"
        session["_fresh"] = True
        session["current_feid"] = 3
    assert client.get("/dashboard/transfers").status_code == 403
    with client.session_transaction() as session:
        session["current_feid"] = 2
    with patch("blueprints.transfer_dashboard.render_template", return_value="dashboard"):
        response = client.get("/dashboard/transfers")
    assert response.status_code == 200


def test_frontend_uses_safe_text_and_mobile_card_layout():
    project = Path(__file__).resolve().parents[1]
    javascript = (project / "static" / "js" / "transfer_dashboard.js").read_text()
    stylesheet = (project / "static" / "css" / "transfer_dashboard.css").read_text()
    template = (project / "templates" / "transfer_dashboard.html").read_text()

    assert "textContent" in javascript
    assert "innerHTML" not in javascript
    assert "td::before" in stylesheet
    assert "grid-template-columns: repeat(auto-fit, minmax(7.6rem, 1fr));" in stylesheet
    assert ".transfer-contextbar { align-items: stretch; flex-direction: column;" in stylesheet
    assert 'data-export-url' in template
    assert 'data-kpi="departures_today"' in template
