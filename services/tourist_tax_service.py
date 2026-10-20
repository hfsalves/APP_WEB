"""Read-only tourist-tax reporting for GuestSpaTur reservations."""

from __future__ import annotations

import calendar
import io
import json
import re
import uuid
from collections import defaultdict
from datetime import date, datetime, timedelta
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import text

from models import db


DELIVERY_MONTHLY = "Mensal"
DELIVERY_QUARTERLY = "Trimestral"
DEFAULT_NIGHT_LIMIT = 7
AMARANTE_NIGHT_LIMIT = 3


def _clean(value: Any) -> str:
    return str(value or "").strip()


def _normalise(value: Any) -> str:
    return " ".join(_clean(value).upper().split())


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not value:
        return None
    try:
        return date.fromisoformat(str(value)[:10])
    except (TypeError, ValueError):
        return None


def _delivery_type(value: Any) -> str:
    return DELIVERY_MONTHLY if _normalise(value) == "MENSAL" else DELIVERY_QUARTERLY


def period_bounds(year: Any, period_type: Any, period_value: Any) -> tuple[date, date, str]:
    try:
        selected_year = int(year)
    except (TypeError, ValueError) as exc:
        raise ValueError("Ano inválido.") from exc
    if not 2000 <= selected_year <= 2100:
        raise ValueError("Ano inválido.")
    kind = _clean(period_type).lower()
    try:
        value = int(period_value)
    except (TypeError, ValueError) as exc:
        raise ValueError("Período inválido.") from exc
    if kind == "mensal":
        if not 1 <= value <= 12:
            raise ValueError("Mês inválido.")
        start = date(selected_year, value, 1)
        return start, date(selected_year, value, calendar.monthrange(selected_year, value)[1]), start.strftime("%Y-%m")
    if kind == "trimestral":
        if not 1 <= value <= 4:
            raise ValueError("Trimestre inválido.")
        start_month = ((value - 1) * 3) + 1
        return date(selected_year, start_month, 1), date(selected_year, start_month + 2, calendar.monthrange(selected_year, start_month + 2)[1]), f"{selected_year}-Q{value}"
    raise ValueError("Tipo de período inválido.")


def reservation_origin(reservation: Any) -> str | None:
    code = _clean(reservation).upper()
    if code.startswith("I"):
        return "Interna"
    if code.startswith("H"):
        return "Airbnb"
    if code[:1].isdigit():
        return "Booking"
    return None


def night_limit(accommodation: Any, zone: Any = "", locality: Any = "") -> int:
    location = " ".join((_normalise(accommodation), _normalise(zone), _normalise(locality)))
    return AMARANTE_NIGHT_LIMIT if "AMARANTE" in location or _normalise(accommodation) == "CASA DO CRASTO" else DEFAULT_NIGHT_LIMIT


def build_tourist_tax_rows(
    rows: Iterable[dict[str, Any]],
    start: date,
    end: date,
    delivery_type: str | None = None,
) -> tuple[list[dict[str, Any]], list[str]]:
    """Apply the ticket's reporting rules to raw RS/AL records."""
    result: list[dict[str, Any]] = []
    warnings: list[str] = []
    for raw in rows:
        accommodation = _clean(raw.get("ALOJAMENTO"))
        mapped = bool(raw.get("ALSTAMP"))
        if not mapped:
            warnings.append(f"Alojamento sem correspondência em AL: {accommodation or '(vazio)'} ({_clean(raw.get('RESERVA')) or 'sem reserva'}).")
            continue
        accommodation_delivery = _delivery_type(raw.get("TAXATURPERIODO"))
        if delivery_type and accommodation_delivery != _delivery_type(delivery_type):
            continue
        origin = reservation_origin(raw.get("RESERVA"))
        if origin == "Interna":
            continue
        if origin is None:
            warnings.append(f"Reserva excluída por prefixo não reconhecido: {_clean(raw.get('RESERVA')) or '(vazia)'}.")
            continue
        checkin, checkout = _as_date(raw.get("DATAIN")), _as_date(raw.get("DATAOUT"))
        try:
            official_nights = int(raw.get("NOITES") or 0)
        except (TypeError, ValueError):
            official_nights = 0
        if not checkin or not checkout or checkout <= checkin or official_nights <= 0:
            warnings.append(f"Reserva excluída por datas/NOITES inválidas: {_clean(raw.get('RESERVA')) or '(sem código)'}.")
            continue
        real_nights = (checkout - checkin).days
        if real_nights != official_nights:
            warnings.append(f"NOITES incompatível em {_clean(raw.get('RESERVA'))}: oficial {official_nights}, datas {real_nights}; usado o valor oficial.")
        official_checkout = min(checkout, checkin + timedelta(days=official_nights))
        allocated_start, allocated_end = max(checkin, start), min(official_checkout, end + timedelta(days=1))
        nights_in_period = max(0, (allocated_end - allocated_start).days)
        if not nights_in_period:
            continue
        adults = max(0, int(raw.get("ADULTOS") or 0))
        children = max(0, int(raw.get("CRIANCAS") or 0))
        guests = adults + children
        limit = night_limit(accommodation, raw.get("ZONA"), raw.get("LOCAL"))
        first_night_index = (allocated_start - checkin).days
        nights_above_limit = sum(1 for index in range(first_night_index, first_night_index + nights_in_period) if index >= limit)
        result.append({
            "reserva": _clean(raw.get("RESERVA")), "origem": origin, "alojamento": accommodation,
            "checkin": checkin, "checkout": checkout, "noites_oficiais": official_nights,
            "noites_periodo": nights_in_period, "adultos": adults, "criancas": children,
            "hospedes_total": guests, "dormidas_totais": guests * nights_in_period,
            "dormidas_acima_limite": guests * nights_above_limit, "limite_noites": limit,
            "entrega": accommodation_delivery,
        })
    return sorted(result, key=lambda row: (row["alojamento"], row["checkin"], row["reserva"])), warnings


def _summary(rows: Iterable[dict[str, Any]], key: str) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, Any]] = defaultdict(lambda: {"reservas": 0, "dormidas_totais": 0, "dormidas_acima_limite": 0, "adultos": 0, "criancas": 0, "noites_periodo": 0})
    seen: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        bucket = grouped[row[key]]
        reservation = row["reserva"]
        if reservation not in seen[row[key]]:
            bucket["reservas"] += 1
            bucket["adultos"] += row["adultos"]
            bucket["criancas"] += row["criancas"]
            seen[row[key]].add(reservation)
        for field in ("dormidas_totais", "dormidas_acima_limite", "noites_periodo"):
            bucket[field] += row[field]
    return [{key: name, **values} for name, values in sorted(grouped.items())]


def tourist_tax_report(year: Any, period_type: Any, period_value: Any, feid: int = 2) -> dict[str, Any]:
    start, end, label = period_bounds(year, period_type, period_value)
    selected_delivery = DELIVERY_MONTHLY if _clean(period_type).lower() == "mensal" else DELIVERY_QUARTERLY
    raw_rows = db.session.execute(text("""
        SELECT RS.RESERVA, RS.ALOJAMENTO, CAST(RS.DATAIN AS date) AS DATAIN,
               CAST(RS.DATAOUT AS date) AS DATAOUT, RS.NOITES, RS.ADULTOS, RS.CRIANCAS,
               AL.ALSTAMP, AL.ZONA, AL.LOCAL,
               ISNULL(NULLIF(LTRIM(RTRIM(AL.TAXATURPERIODO)), ''), 'Trimestral') AS TAXATURPERIODO
        FROM dbo.RS RS
        LEFT JOIN dbo.AL AL
          ON LTRIM(RTRIM(ISNULL(AL.NOME, ''))) COLLATE SQL_Latin1_General_CP1_CI_AI
           = LTRIM(RTRIM(ISNULL(RS.ALOJAMENTO, ''))) COLLATE SQL_Latin1_General_CP1_CI_AI
         AND ISNULL(AL.FEID, 0) = :feid
        WHERE ISNULL(RS.FEID, 0) = :feid
          AND ISNULL(RS.CANCELADA, 0) = 0
          AND CAST(RS.DATAIN AS date) < DATEADD(day, 1, :end_date)
          AND CAST(RS.DATAOUT AS date) > :start_date
    """), {"feid": int(feid), "start_date": start, "end_date": end}).mappings().all()
    rows, warnings = build_tourist_tax_rows(
        [dict(row) for row in raw_rows], start, end, delivery_type=selected_delivery
    )
    accommodation_summary = _summary(rows, "alojamento")
    delivery_by_accommodation = {row["alojamento"]: row["entrega"] for row in rows}
    for item in accommodation_summary:
        item["entrega"] = delivery_by_accommodation.get(item["alojamento"], selected_delivery)
    origin_summary = _summary(rows, "origem")
    configured_rows = db.session.execute(text("""
        SELECT
            CASE
                WHEN UPPER(LTRIM(RTRIM(ISNULL(TAXATURPERIODO, '')))) = 'MENSAL' THEN 'Mensal'
                ELSE 'Trimestral'
            END AS TIPO,
            COUNT(*) AS TOTAL
        FROM dbo.AL
        WHERE ISNULL(FEID, 0) = :feid
          AND ISNULL(INATIVO, 0) = 0
        GROUP BY CASE
            WHEN UPPER(LTRIM(RTRIM(ISNULL(TAXATURPERIODO, '')))) = 'MENSAL' THEN 'Mensal'
            ELSE 'Trimestral'
        END
    """), {"feid": int(feid)}).mappings().all()
    configured_counts = {str(item["TIPO"]): int(item["TOTAL"] or 0) for item in configured_rows}
    return {
        "periodo": {
            "inicio": start.isoformat(), "fim": end.isoformat(), "rotulo": label,
            "tipo": selected_delivery, "feid": int(feid),
        },
        "linhas": rows,
        "resumo_alojamento": accommodation_summary,
        "resumo_origem": origin_summary,
        "totais": {
            "reservas": sum(item["reservas"] for item in accommodation_summary),
            "alojamentos": len(accommodation_summary),
            "dormidas_totais": sum(item["dormidas_totais"] for item in accommodation_summary),
            "dormidas_acima_limite": sum(item["dormidas_acima_limite"] for item in accommodation_summary),
        },
        "configuracao": {
            "tipo_selecionado": selected_delivery,
            "alojamentos_mensais": configured_counts.get(DELIVERY_MONTHLY, 0),
            "alojamentos_trimestrais": configured_counts.get(DELIVERY_QUARTERLY, 0),
        },
        "avisos": warnings,
    }


def build_tourist_tax_workbook(report: dict[str, Any]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill, header_font = PatternFill("solid", fgColor="1F4E78"), Font(color="FFFFFF", bold=True)
    used_sheet_names: set[str] = set()

    def add_sheet(name: str, headers: list[str], data: list[list[Any]]) -> None:
        base_name = re.sub(r"[\\/*?:\[\]]", "-", _clean(name))[:31].strip() or "Mapa"
        sheet_name, suffix = base_name, 2
        while sheet_name.lower() in used_sheet_names:
            marker = f" ({suffix})"
            sheet_name = f"{base_name[:31 - len(marker)]}{marker}"
            suffix += 1
        used_sheet_names.add(sheet_name.lower())
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(headers)
        for cell in sheet[1]: cell.fill, cell.font = header_fill, header_font
        for item in data: sheet.append(item)
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
        date_columns = {index + 1 for index, header in enumerate(headers) if header in {"Check-in", "Check-out", "Data"}}
        for row in sheet.iter_rows(min_row=2):
            for column_index in date_columns:
                row[column_index - 1].number_format = "dd/mm/yyyy"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(36, max(12, max(len(str(cell.value or "")) for cell in column) + 2))

    rows = report["linhas"]
    detail_headers = ["N.º reserva", "Origem", "Check-in", "Check-out", "Noites no período", "Hóspedes adultos", "Hóspedes crianças", "Hóspedes total", "Dormidas totais", "Dormidas acima limite", "Limite noites"]
    def detail(row): return [row["reserva"], row["origem"], row["checkin"], row["checkout"], row["noites_periodo"], row["adultos"], row["criancas"], row["hospedes_total"], row["dormidas_totais"], row["dormidas_acima_limite"], row["limite_noites"]]
    summary_rows = [[row["alojamento"], row["entrega"], row["reservas"], row["dormidas_totais"], row["dormidas_acima_limite"], row["adultos"], row["criancas"], row["noites_periodo"]] for row in report["resumo_alojamento"]]
    if summary_rows: summary_rows.append(["Total", "", *[sum(row[index] for row in summary_rows) for index in range(2, 8)]])
    add_sheet("Resumo", ["Alojamento", "Entrega", "Reservas", "Dormidas totais", "Dormidas acima limite", "Hóspedes adultos", "Hóspedes crianças", "Noites no período"], summary_rows)
    for accommodation in sorted({row["alojamento"] for row in rows}):
        accommodation_rows = [detail(row) for row in rows if row["alojamento"] == accommodation]
        if accommodation_rows: accommodation_rows.append(["Total", "", "", "", sum(row[4] for row in accommodation_rows), sum(row[5] for row in accommodation_rows), sum(row[6] for row in accommodation_rows), sum(row[7] for row in accommodation_rows), sum(row[8] for row in accommodation_rows), sum(row[9] for row in accommodation_rows), ""])
        add_sheet(accommodation, detail_headers, accommodation_rows)
    detail_rows = [detail(row) for row in rows]
    if detail_rows: detail_rows.append(["Total", "", "", "", sum(row[4] for row in detail_rows), sum(row[5] for row in detail_rows), sum(row[6] for row in detail_rows), sum(row[7] for row in detail_rows), sum(row[8] for row in detail_rows), sum(row[9] for row in detail_rows), ""])
    add_sheet("Detalhe", detail_headers, detail_rows)
    origin_rows = [[row["origem"], row["reservas"], row["dormidas_totais"], row["dormidas_acima_limite"], row["adultos"], row["criancas"], row["noites_periodo"]] for row in report["resumo_origem"]]
    if origin_rows: origin_rows.append(["Total", *[sum(row[index] for row in origin_rows) for index in range(1, 7)]])
    add_sheet("Resumo por origem", ["Origem", "Reservas", "Dormidas totais", "Dormidas acima limite", "Hóspedes adultos", "Hóspedes crianças", "Noites no período"], origin_rows)
    add_sheet("Notas e metodologia", ["Nota"], [[f"Período: {report['periodo']['rotulo']} ({report['periodo']['tipo']})"], ["O mapa inclui apenas os alojamentos configurados no AL com esta periodicidade de entrega."], ["Origem: H = Airbnb; prefixo numérico = Booking; I e formatos desconhecidos são excluídos."], ["As noites usam a interseção do período com NOITES oficial; o limite é incremental desde o check-in."], *[[warning] for warning in report["avisos"]]])
    output = io.BytesIO(); workbook.save(output); return output.getvalue()


def audit_tourist_tax_access(user_login: str, report: dict[str, Any], exported: bool = False) -> None:
    """Record only the user, period and action—never credentials or reservation data."""
    db.session.execute(text("""
        IF OBJECT_ID('dbo.GUESTSPA_TOURIST_TAX_AUDIT', 'U') IS NULL
        CREATE TABLE dbo.GUESTSPA_TOURIST_TAX_AUDIT (
            AUDITSTAMP varchar(25) NOT NULL PRIMARY KEY, UTILIZADOR varchar(60) NOT NULL DEFAULT '',
            PERIODO varchar(20) NOT NULL DEFAULT '', TIPO varchar(20) NOT NULL DEFAULT '',
            DATA_INI date NULL, DATA_FIM date NULL, FEID int NULL,
            ACAO varchar(20) NOT NULL DEFAULT '', FILTROS nvarchar(500) NOT NULL DEFAULT N'',
            DTCRIACAO_UTC datetime2 NOT NULL DEFAULT SYSUTCDATETIME()
        )
    """))
    for column_name, definition in (
        ("TIPO", "varchar(20) NOT NULL CONSTRAINT DF_GSTTA_TIPO DEFAULT ('') WITH VALUES"),
        ("DATA_INI", "date NULL"), ("DATA_FIM", "date NULL"), ("FEID", "int NULL"),
        ("FILTROS", "nvarchar(500) NOT NULL CONSTRAINT DF_GSTTA_FILTROS DEFAULT (N'') WITH VALUES"),
    ):
        if not db.session.execute(text("SELECT COL_LENGTH('dbo.GUESTSPA_TOURIST_TAX_AUDIT', :column_name)"), {"column_name": column_name}).scalar():
            db.session.execute(text(f"ALTER TABLE dbo.GUESTSPA_TOURIST_TAX_AUDIT ADD {column_name} {definition}"))
    period = report["periodo"]
    filters = json.dumps({"tipo": period["tipo"], "inicio": period["inicio"], "fim": period["fim"]}, ensure_ascii=False)
    db.session.execute(text("""
        INSERT INTO dbo.GUESTSPA_TOURIST_TAX_AUDIT
            (AUDITSTAMP, UTILIZADOR, PERIODO, TIPO, DATA_INI, DATA_FIM, FEID, ACAO, FILTROS)
        VALUES
            (:stamp, :user, :period, :type, :start, :end, :feid, :action, :filters)
    """), {
        "stamp": uuid.uuid4().hex.upper()[:25], "user": _clean(user_login)[:60],
        "period": period["rotulo"], "type": period["tipo"], "start": period["inicio"],
        "end": period["fim"], "feid": period["feid"],
        "action": "EXPORTAR" if exported else "CONSULTAR", "filters": filters,
    })
    db.session.commit()
