"""Read-only operational dashboard for GuestSpaTur transfer requests."""

from __future__ import annotations

import io
from collections import Counter
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from typing import Any, Iterable, Mapping
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import inspect, text


LISBON = ZoneInfo("Europe/Lisbon")
TARGET_FEID = 2
MAX_CUSTOM_RANGE_DAYS = 731
VALID_PERIODS = frozenset({"today", "next7", "custom", "all"})
VALID_MOVEMENTS = frozenset({"", "arrival", "departure"})
ROUTE_LABELS = {
    "airport": "Aeroporto do Porto → alojamento",
    "campanha": "Campanhã → alojamento",
    "to_airport": "Alojamento → Aeroporto do Porto",
    "to_campanha": "Alojamento → Campanhã",
}
STATUS_LABELS = {
    "PENDENTE": "Pendente",
    "ENVIADO": "Enviado",
    "CANCELADO": "Cancelado",
    "CANCELADA": "Cancelado",
    "EMAIL_ERRO": "Erro no email",
}


class TransferDashboardError(ValueError):
    pass


def _clean(value: Any) -> str:
    return str(value or "").strip()


def _normal(value: Any) -> str:
    return " ".join(_clean(value).upper().split())


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = _clean(value)
    if not raw:
        return None
    try:
        return date.fromisoformat(raw[:10])
    except ValueError:
        return None


def _as_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    raw = _clean(value)
    if not raw:
        return None
    try:
        return datetime.fromisoformat(raw.replace("Z", "+00:00"))
    except ValueError:
        return None


def _as_int(value: Any) -> int:
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _as_money(value: Any) -> float:
    try:
        return float(Decimal(str(value or 0)).quantize(Decimal("0.01")))
    except (InvalidOperation, TypeError, ValueError):
        return 0.0


def _parse_date(value: Any, label: str) -> date:
    parsed = _as_date(value)
    if not parsed:
        raise TransferDashboardError(f"Indique uma {label} válida.")
    return parsed


def transfer_period(mode: Any = None, start: Any = None, end: Any = None, *, today: date | None = None) -> dict[str, Any]:
    today = today or datetime.now(LISBON).date()
    selected = _clean(mode).lower() or "next7"
    if selected not in VALID_PERIODS:
        raise TransferDashboardError("Período inválido.")
    if selected == "all":
        return {"mode": selected, "start": None, "end": None, "label": "Todos os pedidos"}
    if selected == "today":
        start_date = end_date = today
        label = "Hoje"
    elif selected == "next7":
        start_date, end_date = today, today + timedelta(days=7)
        label = "Hoje e próximos 7 dias"
    else:
        start_date = _parse_date(start, "data inicial")
        end_date = _parse_date(end, "data final")
        label = f"{start_date:%d/%m/%Y} a {end_date:%d/%m/%Y}"
    if end_date < start_date:
        raise TransferDashboardError("A data final não pode ser anterior à inicial.")
    if (end_date - start_date).days + 1 > MAX_CUSTOM_RANGE_DAYS:
        raise TransferDashboardError(f"O período máximo é de {MAX_CUSTOM_RANGE_DAYS} dias.")
    return {"mode": selected, "start": start_date, "end": end_date, "label": label}


def transfer_filters(values: Mapping[str, Any] | None = None) -> dict[str, str]:
    values = values or {}
    movement = _clean(values.get("movement")).lower()
    if movement not in VALID_MOVEMENTS:
        movement = ""
    return {
        "accommodation": _clean(values.get("accommodation"))[:120],
        "origin": _clean(values.get("origin")).lower()[:20],
        "status": _normal(values.get("status"))[:20],
        "reservation": _clean(values.get("reservation"))[:60],
        "movement": movement,
    }


def _combine_movement(day: date | None, hour: Any) -> datetime | None:
    if not day:
        return None
    raw = _clean(hour)
    try:
        parsed_time = time.fromisoformat(raw) if raw else time.min
    except ValueError:
        parsed_time = time.min
    return datetime.combine(day, parsed_time)


def _serialize(raw: Mapping[str, Any]) -> dict[str, Any]:
    origin = _clean(raw.get("ORIGEM")).lower()
    is_departure_route = origin.startswith("to_")
    arrival_date = _as_date(raw.get("DATA_CHEGADA"))
    departure_date = _as_date(raw.get("DATA_PARTIDA"))
    has_arrival = not is_departure_route and bool(arrival_date)
    has_departure = bool(raw.get("TEM_PARTIDA")) and bool(departure_date)
    arrival_at = _combine_movement(arrival_date, raw.get("HORA_CHEGADA")) if has_arrival else None
    departure_at = _combine_movement(departure_date, raw.get("HORA_PARTIDA")) if has_departure else None
    status = _normal(raw.get("ESTADO")) or "PENDENTE"
    guest_name = _clean(raw.get("NOME_HOSPEDE") or raw.get("RS_NOME"))
    people = _as_int(raw.get("NUM_PESSOAS"))
    if people <= 0:
        people = _as_int(raw.get("RS_ADULTOS")) + _as_int(raw.get("RS_CRIANCAS")) + _as_int(raw.get("RS_BEBES"))
    incomplete = []
    if not guest_name:
        incomplete.append("hóspede")
    if people <= 0:
        incomplete.append("pessoas")
    if not origin:
        incomplete.append("percurso")
    if not arrival_at and not departure_at:
        incomplete.append("data/hora")
    created = _as_datetime(raw.get("DTCRI"))
    updated = _as_datetime(raw.get("DTALT"))
    return {
        "transfer_stamp": _clean(raw.get("TRANSFERSTAMP")),
        "rsstamp": _clean(raw.get("RSSTAMP")),
        "reservation": _clean(raw.get("RESERVA") or raw.get("RS_RESERVA")),
        "accommodation": _clean(raw.get("ALOJAMENTO") or raw.get("RS_ALOJAMENTO")),
        "origin": origin,
        "origin_label": ROUTE_LABELS.get(origin, origin or "Sem percurso"),
        "guest_name": guest_name,
        "people": people,
        "arrival_date": arrival_date.isoformat() if has_arrival and arrival_date else None,
        "arrival_time": _clean(raw.get("HORA_CHEGADA")) if has_arrival else "",
        "arrival_at": arrival_at.isoformat(timespec="minutes") if arrival_at else None,
        "flight": _clean(raw.get("NUM_VOO")),
        "has_departure": has_departure,
        "departure_date": departure_date.isoformat() if departure_date else None,
        "departure_time": _clean(raw.get("HORA_PARTIDA")) if has_departure else "",
        "departure_at": departure_at.isoformat(timespec="minutes") if departure_at else None,
        "value": _as_money(raw.get("VALOR")),
        "status": status,
        "status_label": STATUS_LABELS.get(status, status.title()),
        "created_at": created.isoformat(timespec="seconds") if created else None,
        "updated_at": updated.isoformat(timespec="seconds") if updated else None,
        "incomplete": bool(incomplete),
        "incomplete_fields": incomplete,
    }


def _date_in_period(value: Any, period: Mapping[str, Any]) -> bool:
    if period.get("mode") == "all":
        return bool(value)
    parsed = _as_date(value)
    return bool(parsed and period["start"] <= parsed <= period["end"])


def _matches_period(row: Mapping[str, Any], period: Mapping[str, Any], movement: str) -> bool:
    arrival_matches = _date_in_period(row.get("arrival_date"), period)
    departure_matches = _date_in_period(row.get("departure_date"), period)
    if movement == "arrival":
        return arrival_matches
    if movement == "departure":
        return departure_matches
    if period.get("mode") == "all":
        return True
    return arrival_matches or departure_matches


def _matches_filters(row: Mapping[str, Any], filters: Mapping[str, str], period: Mapping[str, Any]) -> bool:
    if not _matches_period(row, period, filters["movement"]):
        return False
    if filters["accommodation"] and _normal(row.get("accommodation")) != _normal(filters["accommodation"]):
        return False
    if filters["origin"] and row.get("origin") != filters["origin"]:
        return False
    if filters["status"] and _normal(row.get("status")) != filters["status"]:
        return False
    if filters["reservation"]:
        needle = filters["reservation"].casefold()
        haystack = " ".join((_clean(row.get("reservation")), _clean(row.get("guest_name")), _clean(row.get("transfer_stamp")))).casefold()
        if needle not in haystack:
            return False
    return True


def _sort_key(row: Mapping[str, Any], now: datetime) -> tuple[Any, ...]:
    movements = [_as_datetime(row.get("arrival_at")), _as_datetime(row.get("departure_at"))]
    movements = [movement for movement in movements if movement]
    nearest = min(movements, key=lambda value: abs(value - now)) if movements else datetime.max
    return (nearest, _clean(row.get("reservation")), _clean(row.get("transfer_stamp")))


def _empty_payload(period: Mapping[str, Any], filters: Mapping[str, str]) -> dict[str, Any]:
    return {
        "period": {
            "mode": period["mode"],
            "start": period["start"].isoformat() if period.get("start") else None,
            "end": period["end"].isoformat() if period.get("end") else None,
            "label": period["label"],
        },
        "filters": dict(filters),
        "options": {"accommodations": [], "origins": [], "statuses": []},
        "kpis": {"total": 0, "pending": 0, "sent": 0, "cancelled": 0, "arrivals_today": 0, "departures_today": 0},
        "rows": [],
    }


def _fetch_rows(engine, feid: int) -> list[dict[str, Any]]:
    prefix = "dbo." if engine.dialect.name == "mssql" else ""
    statement = text(f"""
        SELECT
            T.TRANSFERSTAMP, T.RSSTAMP, T.RESERVA, T.ALOJAMENTO, T.ORIGEM,
            T.DATA_CHEGADA, T.HORA_CHEGADA, T.NUM_VOO, T.TEM_PARTIDA,
            T.DATA_PARTIDA, T.HORA_PARTIDA, T.VALOR, T.ESTADO, T.DTCRI, T.DTALT,
            T.NOME_HOSPEDE, T.NUM_PESSOAS,
            R.RESERVA AS RS_RESERVA, R.ALOJAMENTO AS RS_ALOJAMENTO, R.NOME AS RS_NOME,
            R.ADULTOS AS RS_ADULTOS, R.CRIANCAS AS RS_CRIANCAS, R.BEBES AS RS_BEBES
        FROM {prefix}TRANSFERS AS T
        INNER JOIN {prefix}RS AS R ON R.RSSTAMP = T.RSSTAMP
        WHERE COALESCE(R.FEID, 0) = :feid
    """)
    with engine.connect() as connection:
        return [dict(row) for row in connection.execute(statement, {"feid": int(feid)}).mappings()]


def build_transfer_dashboard(
    engine,
    period: Mapping[str, Any],
    filters: Mapping[str, Any] | None = None,
    *,
    feid: int = TARGET_FEID,
    today: date | None = None,
    now: datetime | None = None,
) -> dict[str, Any]:
    filters = transfer_filters(filters)
    today = today or datetime.now(LISBON).date()
    now = now or datetime.now(LISBON).replace(tzinfo=None)
    table_names = {name.upper() for name in inspect(engine).get_table_names()}
    if not {"TRANSFERS", "RS"}.issubset(table_names):
        return _empty_payload(period, filters)

    all_rows = [_serialize(row) for row in _fetch_rows(engine, feid)]
    reservation_counts = Counter(_normal(row["reservation"]) for row in all_rows if row["reservation"])
    rows = [row for row in all_rows if _matches_filters(row, filters, period)]
    for row in rows:
        row["reservation_transfer_count"] = reservation_counts.get(_normal(row["reservation"]), 1)
        row["multiple_for_reservation"] = row["reservation_transfer_count"] > 1
        row["reservation_url"] = f"/reservas/resumo/{row['reservation']}" if row["reservation"] else f"/reservas/rs/{row['rsstamp']}"
    rows.sort(key=lambda row: _sort_key(row, now))

    status_counts = Counter(_normal(row["status"]) for row in rows)
    kpis = {
        "total": len(rows),
        "pending": status_counts.get("PENDENTE", 0),
        "sent": status_counts.get("ENVIADO", 0),
        "cancelled": status_counts.get("CANCELADO", 0) + status_counts.get("CANCELADA", 0),
        "arrivals_today": sum(1 for row in rows if _as_date(row.get("arrival_date")) == today),
        "departures_today": sum(1 for row in rows if _as_date(row.get("departure_date")) == today),
    }
    options = {
        "accommodations": sorted({_clean(row["accommodation"]) for row in all_rows if _clean(row["accommodation"])}, key=str.casefold),
        "origins": [
            {"value": origin, "label": ROUTE_LABELS.get(origin, origin)}
            for origin in sorted({_clean(row["origin"]) for row in all_rows if _clean(row["origin"])}, key=lambda item: ROUTE_LABELS.get(item, item).casefold())
        ],
        "statuses": [
            {"value": status, "label": STATUS_LABELS.get(status, status.title())}
            for status in sorted({_normal(row["status"]) for row in all_rows if _normal(row["status"])}, key=lambda item: STATUS_LABELS.get(item, item).casefold())
        ],
    }
    return {
        "period": {
            "mode": period["mode"],
            "start": period["start"].isoformat() if period.get("start") else None,
            "end": period["end"].isoformat() if period.get("end") else None,
            "label": period["label"],
        },
        "filters": filters,
        "options": options,
        "kpis": kpis,
        "rows": rows,
    }


def _date_time_label(day: Any, hour: Any) -> str:
    parsed = _as_date(day)
    if not parsed:
        return ""
    return f"{parsed:%d/%m/%Y} {_clean(hour)}".strip()


def build_transfer_workbook(payload: Mapping[str, Any], *, exported_at: datetime | None = None) -> bytes:
    exported_at = exported_at or datetime.now(LISBON)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transfers"
    headers = [
        "Reserva", "Alojamento", "Percurso", "Hóspede", "N.º pessoas",
        "Chegada", "Voo", "Tem partida", "Partida", "Valor", "Estado",
        "Criado em", "Última alteração", "Transfer", "Pedidos da reserva",
    ]
    sheet.append(["Resumo de pedidos de transfer"])
    sheet.append([f"Período: {(payload.get('period') or {}).get('label') or 'Todos os pedidos'}"])
    filters = payload.get("filters") or {}
    active_filters = [f"{key}: {value}" for key, value in filters.items() if value]
    sheet.append(["Filtros: " + (" · ".join(active_filters) if active_filters else "Sem filtros adicionais")])
    sheet.append([f"Exportado em: {exported_at:%d/%m/%Y %H:%M}"])
    sheet.append([])
    sheet.append(headers)
    header_row = sheet.max_row
    for row in payload.get("rows") or []:
        sheet.append([
            row.get("reservation") or "",
            row.get("accommodation") or "",
            row.get("origin_label") or "",
            row.get("guest_name") or "",
            int(row.get("people") or 0),
            _date_time_label(row.get("arrival_date"), row.get("arrival_time")),
            row.get("flight") or "",
            "Sim" if row.get("has_departure") else "Não",
            _date_time_label(row.get("departure_date"), row.get("departure_time")),
            float(row.get("value") or 0),
            row.get("status_label") or "",
            row.get("created_at") or "",
            row.get("updated_at") or "",
            row.get("transfer_stamp") or "",
            int(row.get("reservation_transfer_count") or 1),
        ])

    title_fill = PatternFill("solid", fgColor="17324D")
    header_fill = PatternFill("solid", fgColor="2C5F8A")
    sheet["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    sheet["A1"].fill = title_fill
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    for cell in sheet[header_row]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row_cells in sheet.iter_rows(min_row=header_row + 1):
        row_cells[9].number_format = '#,##0.00 "€"'
    widths = [16, 24, 34, 26, 12, 20, 12, 13, 20, 13, 16, 21, 21, 28, 18]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = f"A{header_row + 1}"
    if sheet.max_row >= header_row:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{sheet.max_row}"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()

