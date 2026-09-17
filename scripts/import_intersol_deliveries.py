#!/usr/bin/env python3
"""Import supplier BL/PF chains from PHC.xlsx into INTERSOL.

The operation is dry-run by default.  It creates one supplier delivery note
per external delivery number and one pre-invoice per external invoice number.
Every BI line keeps the PHC lineage BC BI -> BL BI -> PF BI.
"""

from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

import openpyxl
import pyodbc

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app import app
from modules.gr_subcontractor_measurements.service import _new_stamp, _phc_columns, _phc_value
from services.phc_user_import_service import _phc_conn_str


DB = "INTERSOL"
BC_NDOS, BL_NDOS, PF_NDOS = 102, 130, 218
MONEY = Decimal("0.01")
QTY = Decimal("0.0001")
ZERO = Decimal("0")
_columns: dict[str, set[str]] = {}


class ImportValidationError(Exception):
    pass


def clean(value: Any) -> str:
    return str(value or "").strip()


def normalized(value: Any) -> str:
    return "".join(char for char in clean(value).upper() if char.isalnum())


def decimal(value: Any) -> Decimal:
    return Decimal(str(value or 0).replace(",", ".").strip() or "0")


def money(value: Any) -> Decimal:
    return decimal(value).quantize(MONEY, rounding=ROUND_HALF_UP)


def quantity(value: Any) -> Decimal:
    return decimal(value).quantize(QTY)


def doc_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.fromisoformat(clean(value)[:10]).date()


def insert(cursor, table: str, values: dict[str, Any]) -> None:
    columns = _columns.get(table)
    if columns is None:
        columns = _phc_columns(cursor, table)
        _columns[table] = columns
    values = {key: value for key, value in values.items() if key.lower() in columns}
    if not values:
        raise ImportValidationError(f"Sem colunas para inserir em {table}.")
    cursor.execute(
        f"INSERT INTO dbo.{table} ({', '.join(values)}) VALUES ({', '.join('?' for _ in values)})",
        tuple(values.values()),
    )


def rows(cursor, sql: str, params: tuple[Any, ...] = ()) -> list[dict[str, Any]]:
    cursor.execute(sql, params)
    names = [str(item[0]).upper() for item in cursor.description or []]
    return [dict(zip(names, row)) for row in cursor.fetchall()]


def read_xlsx(path: Path) -> list[dict[str, Any]]:
    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = book.active
    headers = [clean(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    required = {"Nº Facture", "Date Facture", "Nº BL", "Date BL", "Référence Produit", "Quant.", "Prix"}
    missing = required - set(headers)
    if missing:
        raise ImportValidationError("Colunas Excel em falta: " + ", ".join(sorted(missing)))
    bc_column = "PHC Nº BC" if "PHC Nº BC" in headers else "BC Nº" if "BC Nº" in headers else ""
    if not bc_column:
        raise ImportValidationError("Coluna BC em falta: esperava 'PHC Nº BC' ou 'BC Nº'.")
    result = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        source = dict(zip(headers, values))
        if not clean(source.get("Nº Facture")):
            continue
        result.append({
            "invoice": clean(source["Nº Facture"]), "invoice_date": doc_date(source["Date Facture"]),
            "delivery": clean(source["Nº BL"]), "delivery_date": doc_date(source["Date BL"]),
            "ref": clean(source["Référence Produit"]), "design": clean(source.get("Produit")),
            "qty": quantity(source["Quant."]), "price": decimal(source["Prix"]),
            "bc": int(decimal(source[bc_column])),
            "origin": clean(source.get("Agence")),
            # HSOLS source files use Chantier as the supplier label; INTERSOL
            # files use Fournisseur.  Keep both to resolve repeated OBRANO.
            "supplier_names": [clean(source.get("Fournisseur")), clean(source.get("Chantier"))],
        })
    if not result:
        raise ImportValidationError("O Excel não contém linhas utilizáveis.")
    return result


def load_source(cursor, bcs: set[int], supplier_names: dict[int, set[str]]) -> tuple[dict[int, dict[str, Any]], list[dict[str, Any]], dict[int, str]]:
    placeholders = ",".join("?" for _ in bcs)
    headers = rows(cursor, f"""
        SELECT B.* FROM dbo.BO B WITH (UPDLOCK,HOLDLOCK)
        WHERE B.NDOS=? AND B.OBRANO IN ({placeholders})
    """, tuple([BC_NDOS, *sorted(bcs)]))
    by_bc: dict[int, dict[str, Any]] = {}
    for number in bcs:
        expected = {normalized(value) for value in supplier_names.get(number, set()) if normalized(value)}
        candidates = [row for row in headers if int(decimal(row["OBRANO"])) == number]
        matched = [row for row in candidates if any(
            candidate in normalized(row.get("NOME"))
            or normalized(row.get("NOME")) in candidate
            or (candidate[:5] and normalized(row.get("NOME")).startswith(candidate[:5]))
            or (candidate[:3] and candidate[:3] in normalized(row.get("NOME")))
            for candidate in expected
        )]
        open_matched = [row for row in matched if not bool(row.get("FECHADA"))]
        if len(open_matched) == 1:
            matched = open_matched
        if len(matched) != 1:
            detail = ", ".join(clean(row.get("NOME")) for row in candidates[:5]) or "nenhum"
            raise ImportValidationError(
                f"BC {number}: não foi possível identificar univocamente o fornecedor ({detail})."
            )
        by_bc[number] = matched[0]
    missing = bcs - set(by_bc)
    if missing:
        raise ImportValidationError("BC inexistente na INTERSOL: " + ", ".join(map(str, sorted(missing))))
    # Some approved supplier BC are closed administratively while retaining
    # their source lines.  This import preserves their status and validates
    # the remaining quantities below; it never reopens or recloses a BC.
    source = rows(cursor, f"""
        SELECT I.*, I2.QTTENC FROM dbo.BI I WITH (UPDLOCK,HOLDLOCK)
        LEFT JOIN dbo.BI2 I2 WITH (UPDLOCK,HOLDLOCK) ON I2.BI2STAMP=I.BISTAMP
        WHERE I.BOSTAMP IN ({','.join('?' for _ in by_bc)})
        ORDER BY I.BOSTAMP,I.LORDEM,I.BISTAMP
    """, tuple(row["BOSTAMP"] for row in by_bc.values()))
    source_to_bc = {clean(row["BOSTAMP"]): number for number, row in by_bc.items()}
    for line in source:
        line["bc"] = source_to_bc[clean(line["BOSTAMP"])]
        line["remaining"] = quantity(line.get("QTT")) - quantity(line.get("QTT2"))
    series = rows(cursor, "SELECT NDOS,NMDOS FROM dbo.TS WITH (NOLOCK) WHERE NDOS IN (?,?,?)", (BC_NDOS, BL_NDOS, PF_NDOS))
    names = {int(decimal(row["NDOS"])): clean(row["NMDOS"]) for row in series}
    if set(names) != {BC_NDOS, BL_NDOS, PF_NDOS}:
        raise ImportValidationError("Séries BC/BL/PF não estão configuradas na INTERSOL.")
    return by_bc, source, names


def allocate(excel: list[dict[str, Any]], source: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_bc: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for line in source:
        if clean(line.get("REF")):
            by_bc[int(line["bc"])].append(line)
    result = []
    for row in sorted(excel, key=lambda x: (x["bc"], x["delivery_date"], x["delivery"], x["ref"])):
        remaining = row["qty"]
        candidates = [line for line in by_bc[row["bc"]]
                      if line["remaining"] > ZERO and clean(line["REF"]) == row["ref"]
                      and money(line.get("EDEBITO")) == money(row["price"])]
        for line in candidates:
            if remaining <= ZERO:
                break
            part = min(remaining, line["remaining"])
            line["remaining"] -= part
            remaining -= part
            result.append({"excel": row, "source": line, "qty": part})
        if remaining > ZERO:
            raise ImportValidationError(
                f"Sem saldo BC {row['bc']} para BL {row['delivery']}, artigo {row['ref']}, quantidade {remaining}."
            )
    return result


def assert_no_duplicates(cursor, excel: list[dict[str, Any]]) -> None:
    # The external delivery number is the durable, user-visible id.  FREF is
    # intentionally never used by these imports, including on pre-invoices.
    checks = ((BL_NDOS, "MAQUINA", {row["delivery"] for row in excel}),)
    for ndos, field, values in checks:
        placeholders = ",".join("?" for _ in values)
        found = rows(cursor, f"SELECT OBRANO,{field} FROM dbo.BO WITH (UPDLOCK,HOLDLOCK) WHERE NDOS=? AND LTRIM(RTRIM(ISNULL({field},''))) IN ({placeholders})", tuple([ndos, *sorted(values)]))
        if found:
            raise ImportValidationError(f"Já existem {len(found)} documento(s) {ndos}: " + ", ".join(clean(x[field]) for x in found))


def next_number(cursor, ndos: int, year: int) -> int:
    cursor.execute("SELECT ISNULL(MAX(TRY_CONVERT(int,OBRANO)),0)+1 FROM dbo.BO WITH (UPDLOCK,HOLDLOCK) WHERE NDOS=? AND BOANO=?", (ndos, year))
    return int(cursor.fetchone()[0] or 1)


def audit(initials: str) -> dict[str, Any]:
    now = datetime.now()
    return {"ousrinis": initials, "ousrdata": now, "ousrhora": now.strftime("%H:%M:%S"),
            "usrinis": initials, "usrdata": now, "usrhora": now.strftime("%H:%M:%S")}


def create_header(cursor, *, stamp: str, ndos: int, name: str, number: int, when: date,
                  origin: dict[str, Any], external: str, net: Decimal, tax: Decimal,
                  closed: bool, initials: str, pf_machine: str = "") -> None:
    values = {"bostamp": stamp, "ndos": ndos, "nmdos": name, "obrano": number, "boano": when.year,
              "dataobra": when, "dataopen": when, "datafinal": when, "datafecho": when if closed else date(1900,1,1),
              "no": origin.get("NO"), "estab": origin.get("ESTAB"), "nome": clean(origin.get("NOME"))[:55],
              "ncont": clean(origin.get("NCONT")), "morada": clean(origin.get("MORADA")), "local": clean(origin.get("LOCAL")),
              "codpost": clean(origin.get("CODPOST")), "ccusto": clean(origin.get("CCUSTO")), "moeda": clean(origin.get("MOEDA")) or "EURO",
              "fref": "", "maquina": external[:100] if ndos == BL_NDOS else pf_machine[:100],
              "etotaldeb": net, "totaldeb": _phc_value(net), "etotal": net + tax, "total": _phc_value(net + tax),
              "fechada": 1 if closed else 0, "aprovado": 1 if ndos == PF_NDOS else 0, **audit(initials)}
    insert(cursor, "BO", values)
    insert(cursor, "BO2", {"bo2stamp": stamp, "processo": clean(origin.get("CCUSTO")), "subproc": "", "area": "", "armazem": 1, "anulado": 0, **audit(initials)})
    insert(cursor, "BO3", {"bo3stamp": stamp, "arquivadodigital": 0, **audit(initials)})


def create_origin_line(cursor, *, stamp: str, ndos: int, name: str, number: int,
                       when: date, design: str, parent_header: str, lordem: int,
                       initials: str) -> None:
    """PHC-style separator displayed before lines from a different dossier."""
    line_stamp = _new_stamp()
    insert(cursor, "BI", {
        "bistamp": line_stamp, "bostamp": stamp, "ndos": ndos, "nmdos": name,
        "obrano": number, "boano": when.year, "dataobra": when,
        "ref": "", "design": design[:60], "qtt": ZERO, "qtt2": ZERO,
        "lordem": lordem, "oobostamp": parent_header, "fechada": 0,
        **audit(initials),
    })
    insert(cursor, "BI2", {"bi2stamp": line_stamp, "bostamp": stamp,
                           "origbistamp": "", "qttenc": ZERO, **audit(initials)})


def origin_label(header: dict[str, Any], name: str) -> str:
    when = header.get("DATAOBRA")
    date_text = when.strftime("%d/%m/%Y") if isinstance(when, datetime) else ""
    return f"{name} nº {clean(header.get('OBRANO'))} {date_text}".strip()


def pf_machine(items: list[dict[str, Any]], database: str) -> str:
    origins = {clean(item["excel"].get("origin")) for item in items if clean(item["excel"].get("origin"))}
    if database.upper() == "INTERSOL":
        if len(origins) != 1:
            raise ImportValidationError("Uma PF INTERSOL não pode agregar mais de uma agência.")
        return "INTERSOL-" + next(iter(origins)).upper()
    # The HSOLS source has no agency/origin field.  Do not invent one.
    return ""


def create_tax(cursor, stamp: str, items: list[dict[str, Any]], initials: str) -> None:
    totals: dict[tuple[Decimal, int], Decimal] = defaultdict(lambda: ZERO)
    for item in items:
        totals[(decimal(item["source"].get("IVA")), int(decimal(item["source"].get("TABIVA"))))] += item["qty"] * item["excel"]["price"]
    for (rate, code), net in totals.items():
        vat = money(net * rate / Decimal("100"))
        insert(cursor, "BOT", {"botstamp": _new_stamp(), "bostamp": stamp, "codigo": code, "taxa": rate,
                                "ebaseinc": net, "baseinc": _phc_value(net), "evalor": vat, "valor": _phc_value(vat), **audit(initials)})


def create_line(cursor, *, item: dict[str, Any], stamp: str, ndos: int, name: str, number: int,
                when: date, origin_header: dict[str, Any], parent: str, parent_header: str,
                next_ndos: int, next_name: str, next_number_value: int, initials: str) -> str:
    source, excel = item["source"], item["excel"]
    line_stamp = _new_stamp()
    unit_price, total = excel["price"], item["qty"] * excel["price"]
    copied = {
        key: source.get(key.upper())
        for key in ("ref", "unidade", "iva", "tabiva", "ivaincl", "armazem", "stipo", "familia",
                    "pcusto", "epcusto", "prorc", "desconto", "desc2", "desc3", "desc4", "desc5",
                    "desc6", "lobs", "lobs2")
        if source.get(key.upper()) is not None
    }
    copied.update({"bistamp": line_stamp, "bostamp": stamp, "ndos": ndos, "nmdos": name, "obrano": number, "boano": when.year,
                   "dataobra": when, "dataopen": when, "datafinal": when, "datafecho": when,
                   "ref": clean(source.get("REF")), "design": (excel["design"] or clean(source.get("DESIGN")))[:60],
                   "qtt": item["qty"], "qtt2": ZERO, "pu": _phc_value(unit_price), "debito": _phc_value(unit_price), "edebito": unit_price,
                   "ttdeb": _phc_value(total), "ettdeb": total, "pcusto": _phc_value(unit_price), "epcusto": unit_price, "prorc": _phc_value(unit_price),
                   "no": origin_header.get("NO"), "nome": clean(origin_header.get("NOME"))[:55], "ccusto": clean(source.get("CCUSTO") or origin_header.get("CCUSTO")),
                   "lordem": int(decimal(source.get("LORDEM"))) or 1000, "obistamp": parent, "oobistamp": parent, "oobostamp": parent_header,
                   "ndoc": next_ndos, "nmdoc": next_name, "fno": next_number_value, "fechada": 0, **audit(initials)})
    insert(cursor, "BI", copied)
    insert(cursor, "BI2", {"bi2stamp": line_stamp, "bostamp": stamp, "origbistamp": parent, "qttenc": item["qty"] if ndos == BL_NDOS else ZERO, **audit(initials)})
    return line_stamp


def execute(cursor, allocations: list[dict[str, Any]], headers: dict[int, dict[str, Any]], names: dict[int, str], initials: str) -> dict[str, int]:
    by_delivery: dict[str, list[dict[str, Any]]] = defaultdict(list)
    by_invoice: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in allocations:
        by_delivery[item["excel"]["delivery"]].append(item)
        by_invoice[item["excel"]["invoice"]].append(item)
    bl_lines: dict[int, str] = {}
    bl_headers: dict[str, tuple[str, int]] = {}
    created_headers: list[str] = []
    for delivery, items in sorted(by_delivery.items(), key=lambda pair: (pair[1][0]["excel"]["delivery_date"], pair[0])):
        row, origin = items[0]["excel"], headers[items[0]["source"]["bc"]]
        when = row["delivery_date"]; number = next_number(cursor, BL_NDOS, when.year); stamp = _new_stamp()
        net = money(sum((item["qty"] * item["excel"]["price"] for item in items), ZERO))
        tax = money(sum((item["qty"] * item["excel"]["price"] * decimal(item["source"].get("IVA")) / Decimal("100") for item in items), ZERO))
        create_header(cursor, stamp=stamp, ndos=BL_NDOS, name=names[BL_NDOS], number=number, when=when, origin=origin, external=delivery, net=net, tax=tax, closed=False, initials=initials)
        create_tax(cursor, stamp, items, initials)
        created_headers.append(stamp)
        bl_headers[delivery] = (stamp, number)
        last_origin = ""
        position = 0
        for item in items:
            source_header = headers[item["source"]["bc"]]
            current_origin = clean(source_header["BOSTAMP"])
            if current_origin != last_origin:
                position += 1
                create_origin_line(cursor, stamp=stamp, ndos=BL_NDOS, name=names[BL_NDOS], number=number,
                                   when=when, design=origin_label(source_header, names[BC_NDOS]),
                                   parent_header=current_origin, lordem=position * 1000, initials=initials)
                last_origin = current_origin
            position += 1
            line = create_line(cursor, item=item, stamp=stamp, ndos=BL_NDOS, name=names[BL_NDOS], number=number, when=when, origin_header=origin,
                               parent=clean(item["source"]["BISTAMP"]), parent_header=clean(origin["BOSTAMP"]), next_ndos=0, next_name="", next_number_value=0, initials=initials)
            # INTERSOL's dossier trigger updates the BC quantity satisfied when
            # a child BL line is inserted.  Do not add it again here.
            bl_lines[id(item)] = line
    for invoice, items in sorted(by_invoice.items(), key=lambda pair: (pair[1][0]["excel"]["invoice_date"], pair[0])):
        row, origin = items[0]["excel"], headers[items[0]["source"]["bc"]]
        when = row["invoice_date"]; number = next_number(cursor, PF_NDOS, when.year); stamp = _new_stamp()
        net = money(sum((item["qty"] * item["excel"]["price"] for item in items), ZERO))
        tax = money(sum((item["qty"] * item["excel"]["price"] * decimal(item["source"].get("IVA")) / Decimal("100") for item in items), ZERO))
        create_header(cursor, stamp=stamp, ndos=PF_NDOS, name=names[PF_NDOS], number=number, when=when, origin=origin, external=invoice, net=net, tax=tax, closed=False, initials=initials, pf_machine=pf_machine(items, DB))
        create_tax(cursor, stamp, items, initials)
        created_headers.append(stamp)
        last_bl = ""
        last_bc = ""
        position = 0
        for item in items:
            delivery = item["excel"]["delivery"]
            bl_stamp, bl_number = bl_headers[delivery]
            source_header = headers[item["source"]["bc"]]
            if bl_stamp != last_bl:
                position += 1
                create_origin_line(cursor, stamp=stamp, ndos=PF_NDOS, name=names[PF_NDOS], number=number,
                                   when=when, design=f"{names[BL_NDOS]} nº {bl_number}",
                                   parent_header=bl_stamp, lordem=position * 1000, initials=initials)
                last_bl = bl_stamp
            current_bc = clean(source_header["BOSTAMP"])
            if current_bc != last_bc:
                position += 1
                create_origin_line(cursor, stamp=stamp, ndos=PF_NDOS, name=names[PF_NDOS], number=number,
                                   when=when, design=origin_label(source_header, names[BC_NDOS]),
                                   parent_header="", lordem=position * 1000, initials=initials)
                last_bc = current_bc
            position += 1
            line = create_line(cursor, item=item, stamp=stamp, ndos=PF_NDOS, name=names[PF_NDOS], number=number, when=when, origin_header=origin,
                               parent=bl_lines[id(item)], parent_header=bl_stamp, next_ndos=55, next_name="V/Facture", next_number_value=0, initials=initials)
            cursor.execute("UPDATE dbo.BI SET QTT2=0,FECHADA=0,NDOC=?,NMDOC=?,FNO=?,USRINIS=?,USRDATA=?,USRHORA=? WHERE BISTAMP=?", (PF_NDOS, names[PF_NDOS], number, initials, datetime.now(), datetime.now().strftime('%H:%M:%S'), bl_lines[id(item)]))
    # Imports remain open.  QTT2 is deliberately zero on every imported
    # material line and separator line; source BC state is not forced here.
    placeholders = ",".join("?" for _ in created_headers)
    cursor.execute(f"UPDATE dbo.BO SET FECHADA=0,DATAFECHO=? WHERE BOSTAMP IN ({placeholders})", tuple([date(1900, 1, 1), *created_headers]))
    cursor.execute(f"UPDATE dbo.BI SET QTT2=0,FECHADA=0 WHERE BOSTAMP IN ({placeholders})", tuple(created_headers))
    return {"bl": len(by_delivery), "pf": len(by_invoice), "lines": len(allocations)}


def main() -> int:
    global DB
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--execute", action="store_true")
    parser.add_argument("--user", default="APP")
    parser.add_argument("--database", default=DB)
    parser.add_argument("--server", default="")
    parser.add_argument("--remap-bc", action="append", default=[], metavar="FROM:TO")
    args = parser.parse_args()
    DB = clean(args.database).upper()
    excel = read_xlsx(args.xlsx)
    for mapping in args.remap_bc:
        try:
            old_value, new_value = clean(mapping).split(":", 1)
            old_bc, new_bc = int(old_value), int(new_value)
        except (TypeError, ValueError):
            raise ImportValidationError(f"Remapeamento BC inválido: {mapping!r}. Use FROM:TO.")
        changed = 0
        for row in excel:
            if row["bc"] == old_bc:
                row["bc"] = new_bc
                changed += 1
        if not changed:
            raise ImportValidationError(f"O BC {old_bc} não existe no ficheiro para remapear.")
    with app.app_context():
        connection = pyodbc.connect(_phc_conn_str(DB, clean(args.server)), timeout=45, autocommit=False)
        try:
            cursor = connection.cursor(); cursor.execute("SET TRANSACTION ISOLATION LEVEL SERIALIZABLE")
            supplier_names = defaultdict(set)
            for row in excel:
                supplier_names[row["bc"]].update(row["supplier_names"])
            headers, source, names = load_source(cursor, {row["bc"] for row in excel}, supplier_names)
            assert_no_duplicates(cursor, excel)
            allocations = allocate(excel, source)
            summary = {"bl": len({row["delivery"] for row in excel}), "pf": len({row["invoice"] for row in excel}), "lines": len(allocations)}
            print(f"VALIDAÇÃO OK: {summary['bl']} BL, {summary['pf']} PF, {summary['lines']} linhas.")
            if not args.execute:
                connection.rollback(); print("DRY-RUN concluído. Sem alterações."); return 0
            result = execute(cursor, allocations, headers, names, clean(args.user)[:3].upper() or "APP")
            connection.commit(); print(f"IMPORTAÇÃO CONCLUÍDA: {result['bl']} BL, {result['pf']} PF, {result['lines']} linhas.")
            return 0
        except Exception:
            connection.rollback(); raise
        finally:
            connection.close()


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except ImportValidationError as exc:
        print(f"ERRO DE VALIDAÇÃO: {exc}", file=sys.stderr)
        raise SystemExit(2)
